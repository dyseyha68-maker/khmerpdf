"""Compare OCR engines (and Netra tuning variants) on one file.

    python manage.py ocr_compare scan.pdf --page 1 --truth scan.txt
    python manage.py ocr_compare photo.jpg --dpi 200 300 --netra-conf 0.3 0.4 0.5 --netra-decoder ar blockwise
    python manage.py ocr_compare scan.pdf --truth scan.txt --cloud     # also run enabled Claude/Gemini (uses API credit)

Prints, per run: characters, seconds, and (with --truth) the character error rate.
Cloud engines never run unless --cloud is given.
"""
import itertools
import time
import unicodedata

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError


def _normalize(text):
    return ''.join(unicodedata.normalize('NFC', text or '').split())


def cer(hyp, ref):
    """Character error rate: edit distance / reference length (whitespace ignored)."""
    hyp, ref = _normalize(hyp), _normalize(ref)
    if not ref:
        return None
    prev = list(range(len(ref) + 1))
    for i, h in enumerate(hyp, 1):
        cur = [i]
        for j, r in enumerate(ref, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (h != r)))
        prev = cur
    return prev[-1] / len(ref)


def doc_text(doc):
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells if c.text.strip())
    return '\n'.join(parts)


class Command(BaseCommand):
    help = 'Compare OCR engines (and Netra settings) on one file; optional ground-truth CER.'

    def add_arguments(self, parser):
        parser.add_argument('file', help='PDF or image to OCR')
        parser.add_argument('--page', type=int, default=1, help='Page number for PDFs (default 1)')
        parser.add_argument('--lang', default='khm', choices=['eng', 'khm', 'eng+khm'])
        parser.add_argument('--truth', help='UTF-8 text file with the correct text for that page')
        parser.add_argument('--dpi', type=int, nargs='+', default=[200], help='Render DPI(s) to try')
        parser.add_argument('--netra-conf', type=float, nargs='+', default=[0.4], help='YOLO confidence value(s)')
        parser.add_argument('--netra-decoder', nargs='+', default=['ar'], help='Netra decoder(s), e.g. ar blockwise')
        parser.add_argument('--cloud', action='store_true', help='Also run enabled Claude/Gemini (costs API credit)')
        parser.add_argument('--show', type=int, default=160, help='Characters of output to print per run (0 = none)')

    def handle(self, *args, **o):
        import pytesseract
        from docx import Document
        from apps.pdf.ocr.pages import _iter_document_pages
        from apps.pdf.ocr.engines.tesseract import TesseractEngine
        from apps.pdf.ocr.engines.netra import NetraEngine, _netraocr_page

        if getattr(settings, 'TESSERACT_CMD', None):
            pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD

        truth = None
        if o['truth']:
            with open(o['truth'], encoding='utf-8') as fh:
                truth = fh.read()

        rows = []

        def run(label, engine, doc_page, dpi):
            doc = Document()
            t0 = time.time()
            try:
                chars = engine.run(doc, doc_page, doc_page.convert('L'))
                text, err = doc_text(doc), None
            except Exception as e:
                chars, text, err = 0, '', f'{type(e).__name__}: {str(e)[:80]}'
            rows.append((label, dpi, chars, time.time() - t0, cer(text, truth) if truth else None, err))
            if o['show'] and text:
                self.stdout.write(f'  [{label} @ {dpi} dpi] {text[:o["show"]]!r}')

        for dpi in o['dpi']:
            page = self._render(o['file'], o['page'], dpi, _iter_document_pages)

            run('Tesseract', TesseractEngine(o['lang']), page, dpi)

            if o['lang'] in ('khm', 'eng+khm'):
                for conf, dec in itertools.product(o['netra_conf'], o['netra_decoder']):
                    label = f'Netra conf={conf} dec={dec}'
                    try:
                        from netra_ocr.ocr_engine import KhmerOCRPipeline
                        _netraocr_page._pipeline = KhmerOCRPipeline(detector='yolo', conf=conf, decoder=dec)
                    except Exception as e:
                        rows.append((label, dpi, 0, 0.0, None, f'{type(e).__name__}: {str(e)[:80]}'))
                        break  # not installed / bad option: no point trying the other variants
                    run(label, NetraEngine(o['lang']), page, dpi)

            if o['cloud']:
                from apps.pdf.ocr.pipeline import build_engines
                from apps.pdf.tasks import _get_ocr_settings
                engines, _, _ = build_engines(_get_ocr_settings(), o['lang'], o['lang'])
                for eng in engines:
                    if eng.remote:
                        run(eng.name, eng, page, dpi)

        self._report(rows, truth is not None)

    def _render(self, path, page_no, dpi, iter_pages):
        for seq, total, img in iter_pages(path, page_range_str=str(page_no), max_limit=1, dpi=dpi):
            return img
        raise CommandError('Could not render the requested page.')

    def _report(self, rows, has_truth):
        self.stdout.write('')
        head = f'{"Engine":<30}{"DPI":>5}{"Chars":>7}{"Sec":>7}' + (f'{"CER":>8}' if has_truth else '')
        self.stdout.write(head)
        self.stdout.write('-' * len(head))
        for label, dpi, chars, secs, c, err in rows:
            line = f'{label:<30}{dpi:>5}{chars:>7}{secs:>7.1f}'
            if has_truth:
                line += f'{c * 100:>7.1f}%' if c is not None else f'{"-":>8}'
            self.stdout.write(line + (f'   ERROR: {err}' if err else ''))
        if has_truth:
            self.stdout.write('\nCER = share of reference characters wrong (lower is better; whitespace ignored).')
