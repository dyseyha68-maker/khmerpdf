"""Native Word (.docx) builders for OCR output (structured JSON / markdown / plain)."""
import io
import json
import logging
import os

logger = logging.getLogger(__name__)


def _add_structured_json_to_docx(doc, items):
    """
    Directly converts Claude's structured JSON array of document elements
    (headings, paragraphs with explicit alignments, lists, tables) into
    native Microsoft Word elements with accurate Khmer font styling.
    """
    import re
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn

    ALIGN_MAP = {
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
        'left': WD_ALIGN_PARAGRAPH.LEFT,
        'right': WD_ALIGN_PARAGRAPH.RIGHT,
    }

    def _add_styled_runs(paragraph, text, is_bold=False, font_size=11, font_name_override=None):
        parts = text.split('**')
        chars = 0
        for idx, part in enumerate(parts):
            if not part:
                continue
            run_bold = is_bold or (idx % 2 == 1)
            has_khmer = any('\u1780' <= c <= '\u17FF' or '\u19E0' <= c <= '\u19FF' for c in part)
            if has_khmer:
                font_name = font_name_override if font_name_override else 'Khmer OS Battambang'
            else:
                font_name = 'Calibri'

            run = paragraph.add_run(part)
            run.bold = run_bold
            run.font.name = font_name
            run.font.size = Pt(font_size)

            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn('w:cs'), font_name)
            rFonts.set(qn('w:ascii'), font_name)
            rFonts.set(qn('w:hAnsi'), font_name)
            rFonts.set(qn('w:eastAsia'), font_name)

            chars += len(part)
        return chars

    # ── Sanitization 1: Unpack fake borderless tables that trap narrative paragraphs or have empty padding cells ──
    APPROVAL_KEYWORDS = ('បានឃើញ', 'ប្រធាន', 'អគ្គលេខាធិការ', 'ហត្ថលេខា', 'ក្រឡាបញ្ជី', 'ចៅក្រម', 'អ្នកស្នើសុំ', 'សាមីខ្លួន', 'សាលាដំបូង', 'ក្រសួង', 'ព្រះរាជាណាចក្រ', 'តាងនាម')

    sanitized_items = []
    for it in items:
        if not isinstance(it, dict):
            continue
        if it.get('type') == 'table' and it.get('layout') == 'borderless':
            rows = it.get('rows', [])
            all_cells = [c for r in rows for c in r]
            non_empty_cells = [c for c in all_cells if (c.get('text', '') if isinstance(c, dict) else str(c)).strip()]
            
            # Check if this table is actually trapping regular body paragraphs
            is_narrative_table = False
            for c in non_empty_cells:
                c_txt = c.get('text', '') if isinstance(c, dict) else str(c)
                if len(c_txt) > 50 and not any(k in c_txt for k in APPROVAL_KEYWORDS):
                    if any(c_txt.startswith(p) for p in ('អាស្រ័យ', 'តាមកម្មវត្ថុ', 'ខ្ញុំបាទ', 'សូម', 'ដោយហេតុ', 'យោង')):
                        is_narrative_table = True
                        break

            if is_narrative_table or len(non_empty_cells) == 1:
                # Unpack table cells into sequential paragraphs!
                for c in non_empty_cells:
                    c_txt = c.get('text', '') if isinstance(c, dict) else str(c)
                    c_align = c.get('align', 'justify') if isinstance(c, dict) else 'justify'
                    if c_txt.strip():
                        sanitized_items.append({
                            'type': 'paragraph',
                            'align': c_align,
                            'text': c_txt.strip()
                        })
                continue

            # Check if a top letterhead cell trapped the document title (e.g. \n\nលិខិតអញ្ជើញមេធាវី)
            TITLE_SPLIT_RE = re.compile(r'\n+(លិខិតអញ្ជើញ[^\n]*|សេចក្តីសម្រេច[^\n]*|សាលក្រម[^\n]*|សាលដីកា[^\n]*|ដីកា[^\n]*|ប្រកាស[^\n]*|អនុក្រឹត្យ[^\n]*|ព្រះរាជក្រម[^\n]*|កិច្ចសន្យា[^\n]*|ពាក្យស្នើសុំ[^\n]*|DECISION[^\n]*|PRAKAS[^\n]*)', re.IGNORECASE)
            extracted_title = None
            for r in rows:
                for c in r:
                    c_txt = c.get('text', '') if isinstance(c, dict) else str(c)
                    if any(k in c_txt for k in ('ព្រះរាជាណាចក្រ', 'ជាតិ សាសនា')) and TITLE_SPLIT_RE.search(c_txt):
                        m = TITLE_SPLIT_RE.search(c_txt)
                        extracted_title = m.group(1).strip()
                        new_txt = TITLE_SPLIT_RE.sub('', c_txt).strip()
                        if isinstance(c, dict):
                            c['text'] = new_txt

            sanitized_items.append(it)
            if extracted_title:
                sanitized_items.append({
                    'type': 'heading',
                    'level': 1,
                    'align': 'center',
                    'text': extracted_title
                })
            continue

        sanitized_items.append(it)

    # ── Sanitization 2: Coalesce consecutive tail signature/approval paragraphs into a multi-column borderless table ──
    NARRATIVE_OPENERS = ('អាស្រ័យ', 'តាមកម្មវត្ថុ', 'ខ្ញុំបាទ', 'សូម', 'ដោយហេតុ', 'យោង', 'កម្មវត្ថុ')

    def _is_approval_item(it):
        if not isinstance(it, dict) or it.get('type') != 'paragraph':
            return False
        txt = it.get('text', '').strip()
        if any(txt.startswith(n) for n in NARRATIVE_OPENERS):
            return False
        is_approval = txt.startswith(('បានឃើញ', 'ហត្ថលេខា')) or ('\n' in txt and any(k in txt for k in ('ប្រធាន', 'អគ្គលេខាធិការ', 'ចៅក្រម', 'ក្រឡាបញ្ជី')))
        return is_approval and len(txt) < 160

    processed_items = []
    idx = 0
    total_items = len(sanitized_items)
    while idx < total_items:
        it = sanitized_items[idx]
        if _is_approval_item(it):
            group = [it]
            j = idx + 1
            while j < total_items and _is_approval_item(sanitized_items[j]):
                group.append(sanitized_items[j])
                j += 1
            if len(group) in (2, 3):
                # Build 1-row borderless table
                row_cells = []
                for g_it in group:
                    row_cells.append({
                        'text': g_it.get('text', ''),
                        'align': g_it.get('align', 'center'),
                        'bold': False
                    })
                processed_items.append({
                    'type': 'table',
                    'layout': 'borderless',
                    'rows': [row_cells]
                })
                idx = j
                continue
        processed_items.append(it)
        idx += 1

    total_chars = 0
    for item in processed_items:
        if not isinstance(item, dict):
            continue
        itype = item.get('type', 'paragraph').lower()
        align_str = item.get('align', '').lower()
        text = item.get('text', '')

        if itype == 'heading':
            level = item.get('level', 1)
            size_map = {1: 15.0, 2: 13.0, 3: 12.0}
            font_size = size_map.get(level, 13.0)
            align = ALIGN_MAP.get(align_str, WD_ALIGN_PARAGRAPH.CENTER)
            
            p = doc.add_paragraph()
            p.alignment = align
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_before = Pt(6 if level == 1 else 4)
            p.paragraph_format.space_after = Pt(2)

            total_chars += _add_styled_runs(p, text, is_bold=True, font_size=font_size)

        elif itype == 'list_item':
            is_sub = item.get('sub', False) or bool(re.match(r'^(?:[0-9\u17E0-\u17E9]+\.[0-9\u17E0-\u17E9]+|\([0-9\u17E0-\u17E9a-zA-Z\u1780-\u17B3]+\))', text))
            align = ALIGN_MAP.get(align_str, WD_ALIGN_PARAGRAPH.JUSTIFY if len(text) > 80 else WD_ALIGN_PARAGRAPH.LEFT)

            p = doc.add_paragraph()
            p.alignment = align
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)

            # Hanging indent creates clean bullet/number alignment without Word's duplicate marker bug
            if is_sub:
                p.paragraph_format.left_indent = Inches(0.55)
                p.paragraph_format.first_line_indent = Inches(-0.25)
            else:
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.first_line_indent = Inches(-0.25)

            total_chars += _add_styled_runs(p, text, is_bold=False, font_size=11)

        elif itype == 'table':
            rows_data = item.get('rows', [])
            if not rows_data:
                continue
            layout = item.get('layout', 'grid').lower()
            num_cols = max(len(r) for r in rows_data)
            tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
            tbl.autofit = False
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

            if layout == 'borderless':
                tblPr = tbl._element.xpath('w:tblPr')
                if tblPr:
                    val_none = 'none'
                    borders_xml = f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val_none}"/><w:left w:val="{val_none}"/><w:bottom w:val="{val_none}"/><w:right w:val="{val_none}"/><w:insideH w:val="{val_none}"/><w:insideV w:val="{val_none}"/></w:tblBorders>'
                    tblPr[0].append(parse_xml(borders_xml))
            else:
                tbl.style = 'Table Grid'

            # Calculate balanced column widths for standard printable width (6.5 inches)
            first_row_str = ' '.join(c.get('text', '') if isinstance(c, dict) else str(c) for c in (rows_data[0] if rows_data else []))
            is_top_letterhead = any(k in first_row_str for k in ('សាលាដំបូង', 'ក្រសួង', 'ព្រះរាជាណាចក្រ', 'គណៈមេធាវី', 'តាងនាម'))
            is_bottom_signature = any(k in first_row_str for k in ('ចៅក្រម', 'ប្រធាន', 'មេធាវី', 'ក្រឡាបញ្ជី', '28/026', 'ហត្ថលេខា'))

            if layout == 'borderless' and num_cols == 3:
                if is_top_letterhead:
                    col_widths = [Inches(2.6), Inches(2.5), Inches(1.4)]
                elif is_bottom_signature:
                    col_widths = [Inches(2.2), Inches(1.5), Inches(2.8)]
                else:
                    col_widths = [Inches(2.16), Inches(2.16), Inches(2.16)]
            elif layout == 'borderless' and num_cols == 2:
                if is_top_letterhead:
                    col_widths = [Inches(2.85), Inches(3.65)]
                else:
                    col_widths = [Inches(3.25), Inches(3.25)]
            else:
                col_widths = [Inches(6.5 / num_cols)] * num_cols

            for row_idx, r in enumerate(rows_data):
                row_elem = tbl.rows[row_idx]
                for col_idx in range(num_cols):
                    cell = row_elem.cells[col_idx]
                    if col_idx < len(col_widths):
                        cell.width = col_widths[col_idx]
                    if col_idx >= len(r):
                        continue
                    cell_data = r[col_idx]
                    if isinstance(cell_data, dict):
                        c_text = cell_data.get('text', '')
                        c_align = ALIGN_MAP.get(cell_data.get('align', '').lower(), WD_ALIGN_PARAGRAPH.CENTER)
                        c_bold = cell_data.get('bold', False)
                    else:
                        c_text = str(cell_data)
                        c_align = WD_ALIGN_PARAGRAPH.CENTER
                        c_bold = False

                    lines = [l.strip() for l in re.split(r'\n|<br\s*/?>', c_text) if l.strip()]
                    if not lines:
                        lines = ['']
                    for l_i, l_str in enumerate(lines):
                        p = cell.paragraphs[0] if l_i == 0 else cell.add_paragraph()
                        p.alignment = c_align
                        p.paragraph_format.line_spacing = 1.15
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(2)
                        total_chars += _add_styled_runs(p, l_str, is_bold=c_bold, font_size=10.5)

            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(0)
            sp.paragraph_format.space_after = Pt(2)

        else:  # paragraph
            # Unpack multiline bullet lists (e.g. យោង ៖ \n- ... or ចម្លងជូន ៖ \n- ...)
            if '\n-' in text or '\n•' in text or '\n–' in text:
                sub_lines = [l.strip() for l in text.split('\n') if l.strip()]
                for sl in sub_lines:
                    p = doc.add_paragraph()
                    p.paragraph_format.line_spacing = 1.2
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(2)
                    if sl.startswith(('-', '–', '—', '•', '*', '+')) and not sl.startswith(('--', '---')):
                        if len(sl) > 1 and sl[1] != ' ':
                            sl = sl[0] + ' ' + sl[1:]
                        p.paragraph_format.left_indent = Inches(0.35)
                        p.paragraph_format.first_line_indent = Inches(-0.25)
                        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(sl) > 70 else WD_ALIGN_PARAGRAPH.LEFT
                    elif sl.startswith(('កម្មវត្ថុ', 'យោង', 'ចម្លងជូន')):
                        p.paragraph_format.left_indent = Inches(0.35)
                        p.paragraph_format.first_line_indent = Inches(-0.35)
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    total_chars += _add_styled_runs(p, sl, is_bold=False, font_size=11)
                continue

            # Check if this paragraph is a bullet point or numbered item (with OR WITHOUT space after dash!)
            is_bullet = False
            if text.startswith(('-', '–', '—', '•', '*', '+')) and not text.startswith(('--', '---')):
                is_bullet = True
                if len(text) > 1 and text[1] != ' ':
                    text = text[0] + ' ' + text[1:]
            elif re.match(r'^(?:[0-9\u17E0-\u17E9a-zA-Z\u1780-\u17B3]+[\.\)]|\([0-9\u17E0-\u17E9a-zA-Z\u1780-\u17B3]+\))\s*', text):
                is_bullet = True

            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)

            if is_bullet:
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(text) > 70 else WD_ALIGN_PARAGRAPH.LEFT
            else:
                align = ALIGN_MAP.get(align_str, WD_ALIGN_PARAGRAPH.JUSTIFY if len(text) > 65 else WD_ALIGN_PARAGRAPH.LEFT)
                p.alignment = align

                # Hanging indent for legal labels (e.g. "**កម្មវត្ថុ ៖**" or "**យោង ៖**")
                clean_prefix = re.sub(r'[\*\s]', '', text[:15])
                if clean_prefix.startswith(('កម្មវត្ថុ៖', 'កម្មវត្ថុ:', 'យោង៖', 'យោង:', 'ចម្លងជូន៖', 'ចម្លងជូន:')):
                    p.paragraph_format.left_indent = Inches(0.35)
                    p.paragraph_format.first_line_indent = Inches(-0.35)
                # First line indent for narrative body paragraphs
                elif align == WD_ALIGN_PARAGRAPH.JUSTIFY and len(text) > 70:
                    p.paragraph_format.first_line_indent = Inches(0.4)
                elif text.startswith(('ខ្ញុំបាទ', 'ខ្ញុំនាងខ្ញុំ', 'ខ្ញុំព្រះករុណា', 'តាមកម្មវត្ថុ', 'អាស្រ័យ')):
                    p.paragraph_format.first_line_indent = Inches(0.4)

            total_chars += _add_styled_runs(p, text, is_bold=False, font_size=11)

    return total_chars


def _add_structured_content_to_docx(doc, raw_text):
    """
    Parses Claude's or Gemini's output. If raw_text contains a valid or truncated
    JSON array of structural elements, formats each element with exact alignment,
    styling, and geometry.
    If raw_text is genuine markdown, uses the markdown parser.
    """
    import json
    import logging
    logger = logging.getLogger(__name__)

    clean = raw_text.strip()
    if clean.startswith('```'):
        first_nl = clean.find('\n')
        if first_nl != -1:
            clean = clean[first_nl + 1:]
        if clean.endswith('```'):
            clean = clean[:-3].strip()

    start_idx = clean.find('[')
    if start_idx != -1:
        candidate = clean[start_idx:].strip()

        # Attempt 1: Parse candidate directly with strict=False (allows unescaped newlines in strings)
        try:
            items = json.loads(candidate, strict=False)
            if isinstance(items, list) and len(items) > 0:
                logger.info(f'OCR: Successfully parsed {len(items)} structured JSON elements')
                return _add_structured_json_to_docx(doc, items)
        except Exception as e:
            logger.debug(f'OCR: Direct JSON parse failed: {e}')

        # Attempt 2: Truncation recovery if output was cut off mid-JSON
        last_brace = candidate.rfind('}')
        if last_brace != -1:
            repaired = candidate[:last_brace + 1].rstrip().rstrip(',') + ']'
            try:
                items = json.loads(repaired, strict=False)
                if isinstance(items, list) and len(items) > 0:
                    logger.info(f'OCR: Recovered {len(items)} structured JSON elements from truncated output')
                    return _add_structured_json_to_docx(doc, items)
            except Exception as e:
                logger.warning(f'OCR: Repaired JSON parse failed: {e}')

        # Attempt 3: If candidate looks like JSON code (starts with '[{' or '{\n"type"'),
        # NEVER pass raw JSON code to _add_markdown_to_docx (which dumps JSON braces into Word)!
        # Instead, extract text blocks cleanly.
        if candidate.startswith(('[', '{')) and any(k in candidate for k in ('"type":', '"text":', '"rows":')):
            import re
            extracted_texts = re.findall(r'"text"\s*:\s*"((?:[^"\\]|\\.)*)"', candidate)
            if extracted_texts:
                logger.info(f'OCR: Salvaged {len(extracted_texts)} text blocks from malformed JSON')
                salvaged_chars = 0
                for txt in extracted_texts:
                    clean_t = txt.replace('\\n', '\n').replace('\\"', '"').strip()
                    if clean_t:
                        p = doc.add_paragraph()
                        salvaged_chars += len(clean_t)
                        run = p.add_run(clean_t)
                        run.font.name = 'Khmer OS Battambang'
                if salvaged_chars > 0:
                    return salvaged_chars

    # Fallback to markdown parser only if raw_text is genuine markdown
    return _add_markdown_to_docx(doc, raw_text)


def _add_markdown_to_docx(doc, markdown_text):
    """
    Parse Claude's formatted OCR output (Markdown headers, centered titles,
    bold text, bullet/numbered lists, and tables) into native styled Word
    elements with proper fonts (Khmer OS Battambang for Khmer, Calibri for Latin).
    Returns total extracted character count.
    """
    import re
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn

    def _add_styled_runs(paragraph, text, is_bold=False, font_size=11, font_name_override=None):
        """Parse inline **bold** markers and append runs with explicit Khmer/Latin font switching."""
        parts = text.split('**')
        chars = 0
        for idx, part in enumerate(parts):
            if not part:
                continue
            run_bold = is_bold or (idx % 2 == 1)
            # Detect Khmer characters
            has_khmer = any('\u1780' <= c <= '\u17FF' or '\u19E0' <= c <= '\u19FF' for c in part)
            if has_khmer:
                font_name = font_name_override if font_name_override else 'Khmer OS Battambang'
            else:
                font_name = 'Calibri'

            run = paragraph.add_run(part)
            run.bold = run_bold
            run.font.name = font_name
            run.font.size = Pt(font_size)

            # Explicitly set Word's Complex Script (w:cs) and East Asia fonts
            # so Windows Word renders Khmer characters using Khmer OS Battambang
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn('w:cs'), font_name)
            rFonts.set(qn('w:ascii'), font_name)
            rFonts.set(qn('w:hAnsi'), font_name)
            rFonts.set(qn('w:eastAsia'), font_name)

            chars += len(part)
        return chars

    # Clean markdown code fences if Claude wrapped output in ```markdown ... ```
    text = markdown_text.strip()
    if text.startswith('```'):
        first_newline = text.find('\n')
        if first_newline != -1:
            text = text[first_newline + 1:]
        if text.endswith('```'):
            text = text[:-3].strip()

    lines = text.split('\n')
    i = 0
    n = len(lines)
    total_chars = 0

    while i < n:
        raw_line = lines[i]
        line = raw_line.strip()

        if not line:
            i += 1
            continue

        # ── 1. Check for Markdown Table ──────────────────────────
        if line.startswith('|') and line.endswith('|'):
            table_lines = []
            while i < n and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1

            is_layout_table = any('[LAYOUT]' in tl or '[layout]' in tl for tl in table_lines)

            valid_rows = []
            for r_i, t_line in enumerate(table_lines):
                # If this is a [LAYOUT] table, skip the header column-definition row
                if is_layout_table and r_i == 0:
                    continue
                clean_tline = re.sub(r'\[LAYOUT\]|\[layout\]', '', t_line)
                cells = [c.strip() for c in clean_tline.split('|')[1:-1]]
                if all(set(c).issubset({'-', ':', ' '}) for c in cells if c):
                    continue  # separator row
                valid_rows.append(cells)

            if valid_rows:
                num_cols = max(len(r) for r in valid_rows)
                normalized_rows = [r + [''] * (num_cols - len(r)) for r in valid_rows]
                tbl = doc.add_table(rows=len(normalized_rows), cols=num_cols)
                tbl.autofit = False
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

                has_header_keywords = any(
                    any(k in cell for k in ('ព្រះរាជាណាចក្រ', 'សាលាដំបូង', 'ក្រឡាបញ្ជី', 'ចៅក្រម', 'តាងនាមប្រជារាស្ត្រខ្មែរ'))
                    for row in normalized_rows for cell in row
                )

                if is_layout_table or has_header_keywords:
                    # Apply borderless XML to ensure completely invisible boundaries
                    tblPr = tbl._element.xpath('w:tblPr')
                    if tblPr:
                        val_none = 'none'
                        borders_xml = f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val_none}"/><w:left w:val="{val_none}"/><w:bottom w:val="{val_none}"/><w:right w:val="{val_none}"/><w:insideH w:val="{val_none}"/><w:insideV w:val="{val_none}"/></w:tblBorders>'
                        tblPr[0].append(parse_xml(borders_xml))
                else:
                    tbl.style = 'Table Grid'

                # Explicit column widths for standard printable width (~6.5 inches)
                if is_layout_table or has_header_keywords:
                    if num_cols == 2:
                        col_widths = [Inches(2.85), Inches(3.65)]
                    elif num_cols == 3:
                        col_widths = [Inches(2.16), Inches(2.16), Inches(2.16)]
                    else:
                        col_widths = [Inches(6.5 / num_cols)] * num_cols
                else:
                    col_widths = [Inches(6.5 / num_cols)] * num_cols

                for row_idx, row_cells in enumerate(normalized_rows):
                    is_header = (row_idx == 0 and not is_layout_table and not has_header_keywords)
                    row_elem = tbl.rows[row_idx]
                    for col_idx, cell_text in enumerate(row_cells):
                        cell = row_elem.cells[col_idx]
                        if col_idx < len(col_widths):
                            cell.width = col_widths[col_idx]

                        cell_lines = [cl.strip() for cl in re.split(r'<br\s*/?>', cell_text, flags=re.I) if cl.strip()]
                        if not cell_lines:
                            cell_lines = ['']

                        active_cell_align = None
                        for l_idx, cl in enumerate(cell_lines):
                            if l_idx == 0:
                                p = cell.paragraphs[0]
                            else:
                                p = cell.add_paragraph()

                            if '[CENTER]' in cl or '[center]' in cl:
                                active_cell_align = WD_ALIGN_PARAGRAPH.CENTER
                            elif '[RIGHT]' in cl or '[right]' in cl:
                                active_cell_align = WD_ALIGN_PARAGRAPH.RIGHT
                            elif '[JUSTIFY]' in cl or '[justify]' in cl:
                                active_cell_align = WD_ALIGN_PARAGRAPH.JUSTIFY
                            elif '[LEFT]' in cl or '[left]' in cl:
                                active_cell_align = WD_ALIGN_PARAGRAPH.LEFT

                            if active_cell_align:
                                p_align = active_cell_align
                            elif any(k in cl for k in ('ព្រះរាជាណាចក្រ', 'ជាតិ សាសនា', 'តាងនាមប្រជារាស្ត្រខ្មែរ', 'ក្រឡាបញ្ជី', 'ចៅក្រម', 'ប្រធាន', 'អនុប្រធាន')):
                                p_align = WD_ALIGN_PARAGRAPH.CENTER
                            elif len(cl) > 80:
                                p_align = WD_ALIGN_PARAGRAPH.JUSTIFY
                            else:
                                p_align = WD_ALIGN_PARAGRAPH.LEFT

                            if any(k in cl for k in ('[/CENTER]', '[/center]', '[/RIGHT]', '[/right]', '[/JUSTIFY]', '[/justify]', '[/LEFT]', '[/left]')):
                                active_cell_align = None

                            cl = re.sub(r'\[/?CENTER\]|\[/?center\]|\[/?RIGHT\]|\[/?right\]|\[/?JUSTIFY\]|\[/?justify\]|\[/?LEFT\]|\[/?left\]', '', cl).strip()

                            p.alignment = p_align
                            p.paragraph_format.line_spacing = 1.15
                            p.paragraph_format.space_before = Pt(0)
                            p.paragraph_format.space_after = Pt(2)

                            c_count = _add_styled_runs(p, cl, is_bold=is_header, font_size=10.5 if not is_header else 11)
                            total_chars += c_count

                # Subtle spacing paragraph after table
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(2)
            continue

        # ── 2. Omit standalone page numbers ──────────────────────
        # Standalone '1', '2', '១', '២', '៨' at page margins
        clean_num_check = re.sub(r'\[/?CENTER\]|\[/?RIGHT\]|\[/?center\]|\[/?right\]', '', line).strip()
        if re.match(r'^[0-9\u17E0-\u17E9]{1,3}$', clean_num_check):
            i += 1
            continue

        # ── 3. Alignment & Heading tags ──────────────────────────
        explicit_align = None
        if re.search(r'\[/?CENTER\]|\[/?center\]', line):
            explicit_align = WD_ALIGN_PARAGRAPH.CENTER
            line = re.sub(r'\[/?CENTER\]|\[/?center\]', '', line).strip()
        elif re.search(r'\[/?RIGHT\]|\[/?right\]', line):
            explicit_align = WD_ALIGN_PARAGRAPH.RIGHT
            line = re.sub(r'\[/?RIGHT\]|\[/?right\]', '', line).strip()
        elif re.search(r'\[/?JUSTIFY\]|\[/?justify\]', line):
            explicit_align = WD_ALIGN_PARAGRAPH.JUSTIFY
            line = re.sub(r'\[/?JUSTIFY\]|\[/?justify\]', '', line).strip()
        elif re.search(r'\[/?LEFT\]|\[/?left\]', line):
            explicit_align = WD_ALIGN_PARAGRAPH.LEFT
            line = re.sub(r'\[/?LEFT\]|\[/?left\]', '', line).strip()

        font_size = 11.0
        is_heading_bold = False
        space_before = Pt(0)
        space_after = Pt(3)
        is_heading = False

        # Check Heading levels
        if line.startswith('# '):
            is_heading = True
            font_size = 15.0
            is_heading_bold = True
            space_before = Pt(6)
            space_after = Pt(2)
            line = line[2:].strip()
        elif line.startswith('## '):
            is_heading = True
            font_size = 13.0
            is_heading_bold = True
            space_before = Pt(5)
            space_after = Pt(2)
            line = line[3:].strip()
        elif line.startswith('### '):
            is_heading = True
            font_size = 12.0
            is_heading_bold = True
            space_before = Pt(4)
            space_after = Pt(2)
            line = line[4:].strip()

        # Lists & Sub-clause indentations
        is_bullet = False
        is_number = False
        is_sub = False
        left_indent = None

        bullet_match = re.match(r'^([\u2022\u25CB\u25CF\u25AA\u25AB\u2013\u2014\-\*\+\~])\s+(.+)$', line)
        number_match = re.match(r'^([0-9\u17E0-\u17E9a-zA-Z\u1780-\u17B3]+[\.\)]|\([0-9\u17E0-\u17E9a-zA-Z\u1780-\u17B3]+\))\s+(.+)$', line)
        sub_match = re.match(r'^([0-9\u17E0-\u17E9]+\.[0-9\u17E0-\u17E9]+[\.\)]?)\s+(.+)$', line)

        if sub_match:
            is_sub = True
            is_number = True
        elif number_match:
            is_number = True
        elif bullet_match:
            is_bullet = True
            sym = bullet_match.group(1)
            rest = bullet_match.group(2)
            if sym not in ('-', '–', '—'):
                line = '• ' + rest

        # ── Intelligent Alignment Resolution ──
        # Check if text is center -> make it center; if paragraph/narrative -> make it justify
        if explicit_align is not None:
            align = explicit_align
        elif any(k in line for k in ('ព្រះរាជាណាចក្រកម្ពុជា', 'ជាតិ សាសនា ព្រះមហាក្សត្រ', 'Kingdom of Cambodia', 'Nation Religion King')):
            align = WD_ALIGN_PARAGRAPH.CENTER
        elif is_heading:
            if any(k in line for k in ('សាលក្រម', 'សេចក្តីសម្រេច', 'ដីកា', 'ប្រកាស', 'កិច្ចសន្យា', 'តាងនាមប្រជារាស្ត្រខ្មែរ', 'ព្រះរាជក្រម', 'លិខិត', 'ពាក្យស្នើសុំ', 'NOTICE', 'CERTIFICATE', 'DECISION', 'JUDGMENT', 'CONTRACT', 'ORDER')) or len(line) < 55:
                align = WD_ALIGN_PARAGRAPH.CENTER
            else:
                align = WD_ALIGN_PARAGRAPH.LEFT
        elif re.match(r'^(រាជធានីភ្នំពេញ|ភ្នំពេញ|ធ្វើនៅ|ថ្ងៃទី)\s*,?\s*ថ្ងៃទី', line) or (re.search(r'ថ្ងៃទី\s*[0-9\u17E0-\u17E9]+\s*ខែ\s*[0-9\u17E0-\u17E9\w]+\s*ឆ្នាំ\s*[0-9\u17E0-\u17E9]+', line) and len(line) < 70):
            align = WD_ALIGN_PARAGRAPH.RIGHT
        elif is_bullet or is_number:
            align = WD_ALIGN_PARAGRAPH.JUSTIFY if len(line) > 85 else WD_ALIGN_PARAGRAPH.LEFT
        elif len(line) < 65 or line.endswith(('៖', ':')) or re.match(r'^\*\*[^\*]+(៖|:)\*\*', line):
            align = WD_ALIGN_PARAGRAPH.LEFT
        elif len(line) >= 65:
            align = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            align = WD_ALIGN_PARAGRAPH.LEFT

        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_before = space_before
        p.paragraph_format.space_after = space_after
        if is_sub:
            p.paragraph_format.left_indent = Inches(0.55)
            p.paragraph_format.first_line_indent = Inches(-0.25)
        elif is_bullet or is_number:
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.25)
        else:
            clean_prefix = re.sub(r'[\*\s]', '', line[:15])
            if clean_prefix.startswith(('កម្មវត្ថុ៖', 'កម្មវត្ថុ:', 'យោង៖', 'យោង:')):
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.first_line_indent = Inches(-0.35)
            elif align == WD_ALIGN_PARAGRAPH.JUSTIFY and len(line) > 70:
                p.paragraph_format.first_line_indent = Inches(0.4)
            elif line.startswith(('ខ្ញុំបាទ', 'ខ្ញុំនាងខ្ញុំ', 'ខ្ញុំព្រះករុណា', 'តាមកម្មវត្ថុ', 'អាស្រ័យ')):
                p.paragraph_format.first_line_indent = Inches(0.4)

        c_count = _add_styled_runs(p, line, is_bold=is_heading_bold, font_size=font_size)
        total_chars += c_count
        i += 1

    return total_chars
