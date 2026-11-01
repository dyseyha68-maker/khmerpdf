"""Every OCR Word file: all text 11pt, 1.5-line spacing (tables included); bold/indent/gap still kept."""
from django.test import SimpleTestCase
from docx import Document
from docx.shared import Pt

from apps.pdf.ocr.docx_writer import (
    DOC_FONT_SIZE, DOC_LINE_SPACING, _add_markdown_to_docx, _add_structured_json_to_docx,
)

# elements as the Netra layout step produces them: measured sizes/pitches that must NOT be applied
ITEMS = [
    {'type': 'heading', 'level': 1, 'align': 'center', 'text': 'ចំណងជើង', 'font_size': 18.0, 'bold': True,
     'line_pitch': 32.0},
    {'type': 'paragraph', 'align': 'justify', 'text': 'ខ្លឹមសារ ' * 20, 'font_size': 8.5, 'line_pitch': 15.0,
     'space_before': 12.0, 'first_line_indent': 24},
    {'type': 'list_item', 'text': '1. មាត្រាទី១', 'font_size': 9.0, 'line_pitch': 16.0},
    {'type': 'table', 'layout': 'grid', 'rows': [[{'text': 'ក', 'bold': True}, {'text': 'ខ'}], [{'text': 'គ'}, {'text': 'ឃ'}]]},
    {'type': 'table', 'layout': 'borderless', 'rows': [[{'text': 'ឆ្វេង\nបន្ទាត់២'}, {'text': 'ស្តាំ'}]]},
]


def all_paragraphs(doc):
    paras = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    return paras


class FixedFormatTests(SimpleTestCase):
    def assert_fixed(self, doc):
        paras = [p for p in all_paragraphs(doc) if p.text.strip()]
        self.assertTrue(paras)
        for p in paras:
            self.assertEqual(p.paragraph_format.line_spacing, DOC_LINE_SPACING, p.text[:20])
            for run in p.runs:
                if run.text.strip():
                    self.assertEqual(run.font.size, Pt(DOC_FONT_SIZE), p.text[:20])

    def test_constants_are_what_was_asked(self):
        self.assertEqual((DOC_FONT_SIZE, DOC_LINE_SPACING), (11, 1.5))

    def test_structured_elements_ignore_measured_size_and_pitch(self):
        doc = Document()
        _add_structured_json_to_docx(doc, ITEMS)
        self.assert_fixed(doc)

    def test_tables_use_the_same_size_and_spacing(self):
        doc = Document()
        _add_structured_json_to_docx(doc, ITEMS)
        self.assertEqual(len(doc.tables), 2)
        cells = [c for t in doc.tables for r in t.rows for c in r.cells]
        for cell in cells:
            for p in cell.paragraphs:
                self.assertEqual(p.paragraph_format.line_spacing, 1.5)
                for run in p.runs:
                    self.assertEqual(run.font.size, Pt(11))

    def test_bold_indent_and_gap_are_still_kept(self):
        doc = Document()
        _add_structured_json_to_docx(doc, ITEMS)
        heading, body = doc.paragraphs[0], doc.paragraphs[1]
        self.assertTrue(all(r.bold for r in heading.runs if r.text.strip()))
        self.assertEqual(body.paragraph_format.first_line_indent, Pt(24))
        self.assertEqual(body.paragraph_format.space_before, Pt(12))

    def test_markdown_path_is_fixed_too(self):
        doc = Document()
        _add_markdown_to_docx(doc, '# ចំណងជើងធំ\n\n## ចំណងជើងរង\n\nខ្លឹមសារធម្មតា\n\n| ក | ខ |\n|---|---|\n| គ | ឃ |\n')
        self.assert_fixed(doc)
