"""OCR result preview endpoint: the finished Word file's paragraphs with their formatting."""
import shutil
import tempfile

from django.core.files.base import ContentFile
from django.test import Client, TestCase, override_settings
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt

from apps.pdf.models import Job


def make_docx():
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('Title')
    run.bold = True
    run.font.size = Pt(18)
    body = doc.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.space_before = Pt(12)
    body.paragraph_format.first_line_indent = Pt(24)
    body.paragraph_format.line_spacing = Pt(23)
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    body.add_run('Body text').font.size = Pt(11)
    doc.add_paragraph()                      # empty paragraph: skipped
    doc.add_page_break()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = 'a'
    table.rows[0].cells[1].text = 'b'
    path = tempfile.mktemp(suffix='.docx')
    doc.save(path)
    with open(path, 'rb') as fh:
        return fh.read()


class JobPreviewTests(TestCase):
    def setUp(self):
        self.media = tempfile.mkdtemp()
        self._override = override_settings(MEDIA_ROOT=self.media)
        self._override.enable()

    def tearDown(self):
        self._override.disable()
        shutil.rmtree(self.media, ignore_errors=True)

    def finished_job(self, tool='ocr', status='done'):
        job = Job.objects.create(tool=tool, status=status)
        job.result.save('out.docx', ContentFile(make_docx()))
        return job

    def test_preview_returns_paragraphs_with_formatting(self):
        job = self.finished_job()
        data = Client().get(f'/api/job/{job.id}/preview/').json()
        kinds = [i['type'] for i in data['items']]
        self.assertEqual(kinds, ['p', 'p', 'page_break', 'table'])
        title, body = data['items'][0], data['items'][1]
        self.assertEqual((title['align'], title['bold'], title['size']), ('center', True, 18.0))
        self.assertEqual((body['align'], body['before'], body['indent'], body['pitch']), ('justify', 12.0, 24.0, 23.0))
        self.assertEqual(data['items'][3]['rows'], [['a', 'b']])
        self.assertFalse(data['truncated'])

    def test_no_preview_for_other_tools_or_unfinished_jobs(self):
        self.assertEqual(Client().get(f'/api/job/{self.finished_job(tool="merge").id}/preview/').status_code, 404)
        self.assertEqual(Client().get(f'/api/job/{self.finished_job(status="processing").id}/preview/').status_code, 404)

    def test_failed_pages_are_reported_by_job_status(self):
        job = Job.objects.create(tool='ocr', status='done', files={'current_page': 3, 'total_pages': 3, 'failed_pages': [2]})
        data = Client().get(f'/api/job/{job.id}/').json()
        self.assertEqual(data['failed_pages'], [2])
